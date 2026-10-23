"""OUTLANDS Markdown-to-DOCX builder with explicit input/output paths."""
from __future__ import annotations
import argparse
from datetime import datetime, timezone
import io
import os
from pathlib import Path
import re
import stat
import sys
try:
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:
    _DOCX_AVAILABLE = False
else:
    _DOCX_AVAILABLE = True

ACCENT = "080B11"
ACCENT_DARK = "080B11"
TEAL = "12616B"
LIGHT_ACCENT = "FDFCCF"
LIGHT_TEAL = "E3F8FA"
LIGHT_GRAY = "F1F3F5"
TEXT = "18212F"
MAX_INPUT_BYTES = 262144
MAX_LINES = 5000
MAX_LINE_CHARS = 4000
MAX_COLUMNS = 16
MAX_TABLE_ROWS = 1000

class BuildError(Exception):
    """Bounded input/output failure without host data in its message."""

class _Parser(argparse.ArgumentParser):
    def error(self, message):
        raise BuildError("ARGUMENTS_INVALID")

def split_markdown_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]

def is_table_separator(line: str) -> bool:
    return bool(re.match(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$", line))

def parse_markdown(md: str) -> list[dict]:
    lines = md.replace("\r\n", "\n").split("\n")
    sections: list[dict] = []
    current = {"level": 0, "title": "Front Matter", "lines": []}

    for line in lines:
        match = re.match(r"^(#{1,3})\s+(.+?)\s*$", line)
        if match:
            sections.append(current)
            current = {"level": len(match.group(1)), "title": match.group(2).strip(), "lines": []}
        else:
            current["lines"].append(line)
    sections.append(current)

    for section in sections:
        content = []
        i = 0
        section_lines = section["lines"]
        while i < len(section_lines):
            line = section_lines[i]
            if line.strip().startswith("|") and i + 1 < len(section_lines) and is_table_separator(section_lines[i + 1]):
                rows = [split_markdown_row(line)]
                i += 2
                while i < len(section_lines) and section_lines[i].strip().startswith("|"):
                    rows.append(split_markdown_row(section_lines[i]))
                    i += 1
                content.append({"type": "table", "headers": rows[0], "rows": rows[1:]})
            else:
                content.append({"type": "line", "text": line})
                i += 1
        section["content"] = content

    return [section for section in sections if section["title"] != "Front Matter" or any(line.strip() for line in section["lines"])]

def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)

def set_cell_text(cell, text: str, bold=False, color=TEXT, size=8.4, align=None) -> None:
    cell.text = ""
    para = cell.paragraphs[0]
    if align is not None:
        para.alignment = align
    run = para.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)

def set_paragraph_run_style(run, size=10.2, color=TEXT, bold=False) -> None:
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold

def add_heading(document: Document, text: str, level: int) -> None:
    style = "Heading 1" if level <= 1 else "Heading 2" if level == 2 else "Heading 3"
    para = document.add_paragraph(style=style)
    para.paragraph_format.keep_with_next = True
    run = para.add_run(text)
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor.from_string(ACCENT_DARK if level <= 2 else TEAL)
    run.font.bold = True
    run.font.size = Pt(18 if level <= 1 else 13.5 if level == 2 else 11.2)

def add_callout(document: Document, text: str, fill: str = LIGHT_ACCENT) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=130, start=150, bottom=130, end=150)
    para = cell.paragraphs[0]
    para.paragraph_format.space_after = Pt(0)
    run = para.add_run(text)
    set_paragraph_run_style(run, size=10.5, color=ACCENT_DARK, bold=True)

def add_info_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = table.rows[0].cells
    set_cell_text(headers[0], "Field", bold=True, color="FFFFFF", size=9.2)
    set_cell_text(headers[1], "Details", bold=True, color="FFFFFF", size=9.2)
    set_cell_shading(headers[0], ACCENT)
    set_cell_shading(headers[1], ACCENT)
    for key, value in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], key, bold=True, color=ACCENT_DARK, size=8.8)
        set_cell_text(cells[1], value, color=TEXT, size=8.8)
        set_cell_shading(cells[0], LIGHT_GRAY)
    for row in table.rows:
        row.cells[0].width = Inches(1.75)
        row.cells[1].width = Inches(5.55)

def table_widths(col_count: int) -> list[float]:
    if col_count == 5:
        return [2.05, 1.28, 1.27, 1.32, 1.38]
    if col_count == 4:
        return [2.30, 1.55, 1.70, 1.75]
    if col_count == 3:
        return [2.20, 1.35, 3.75]
    if col_count == 2:
        return [2.25, 5.05]
    return [7.25 / max(col_count, 1)] * col_count

def add_markdown_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    if not headers:
        return
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = table_widths(len(headers))

    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_text(cell, header, bold=True, color="FFFFFF", size=7.8, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(cell, ACCENT)
        cell.width = Inches(widths[idx])
    set_repeat_table_header(table.rows[0])

    for r_idx, row_values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            cell = cells[idx]
            set_cell_text(cell, value, color=TEXT, size=7.6, align=WD_ALIGN_PARAGRAPH.LEFT)
            cell.width = Inches(widths[idx])
            if r_idx % 2 == 1:
                set_cell_shading(cell, "FAFAFC")

    document.add_paragraph().paragraph_format.space_after = Pt(2)

def add_line(document: Document, line: str) -> None:
    text = line.strip()
    if not text:
        return
    if text.startswith("- "):
        para = document.add_paragraph(style="List Bullet")
        run = para.add_run(text[2:].strip())
        set_paragraph_run_style(run, size=9.2)
    elif re.match(r"^\d+\.\s+", text):
        para = document.add_paragraph(style="List Number")
        run = para.add_run(re.sub(r"^\d+\.\s+", "", text))
        set_paragraph_run_style(run, size=9.2)
    else:
        para = document.add_paragraph()
        para.paragraph_format.space_after = Pt(3)
        run = para.add_run(text)
        set_paragraph_run_style(run, size=9.4, bold=text.endswith(":"))

def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.4)
    normal.font.color.rgb = RGBColor.from_string(TEXT)

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor.from_string(ACCENT_DARK)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)


def checked_path(value: str) -> Path:
    path = Path(value)
    if ".." in path.parts:
        raise BuildError("PATH_TRAVERSAL")
    path = path.absolute()
    for item in (path, *path.parents):
        if item.is_symlink():
            raise BuildError("SYMLINK_PATH")
    return path


def read_markdown(path: Path) -> str:
    flags = os.O_RDONLY | getattr(os, "O_NOFOLLOW", 0) | getattr(os, "O_NONBLOCK", 0)
    fd = os.open(path, flags)
    with os.fdopen(fd, "rb") as stream:
        before = os.fstat(stream.fileno())
        if not stat.S_ISREG(before.st_mode):
            raise BuildError("INPUT_NOT_REGULAR")
        if before.st_size > MAX_INPUT_BYTES:
            raise BuildError("INPUT_TOO_LARGE")
        data = stream.read(MAX_INPUT_BYTES + 1)
        after = os.fstat(stream.fileno())
    identity = lambda info: (info.st_dev, info.st_ino, info.st_size, info.st_mtime_ns)
    if len(data) > MAX_INPUT_BYTES:
        raise BuildError("INPUT_TOO_LARGE")
    if identity(before) != identity(after):
        raise BuildError("INPUT_CHANGED")
    value = data.decode("utf-8-sig")
    lines = value.splitlines()
    if len(lines) > MAX_LINES or any(len(line) > MAX_LINE_CHARS for line in lines):
        raise BuildError("INPUT_STRUCTURE_LIMIT")
    if any(ord(char) < 32 and char not in "\n\r\t" for char in value):
        raise BuildError("INPUT_CONTROL_CHARACTER")
    return value


def validate_sections(sections: list[dict]) -> None:
    for section in sections:
        for item in section["content"]:
            if item["type"] != "table":
                continue
            width = len(item["headers"])
            if not 1 <= width <= MAX_COLUMNS or len(item["rows"]) > MAX_TABLE_ROWS:
                raise BuildError("TABLE_LIMIT")
            if any(len(row) != width for row in item["rows"]):
                raise BuildError("TABLE_NOT_RECTANGULAR")


def make_document(markdown: str, title: str) -> bytes:
    sections = parse_markdown(markdown)
    validate_sections(sections)
    document = Document()
    configure_styles(document)
    page = document.sections[0]
    page.top_margin = page.bottom_margin = Inches(0.6)
    page.left_margin = page.right_margin = Inches(0.6)
    add_heading(document, title, 1)
    add_line(document, "OUTLANDS | Generated from the explicitly selected Markdown input.")
    for section in sections:
        if section["title"]:
            add_heading(document, section["title"], max(1, min(section["level"], 3)))
        for item in section["content"]:
            if item["type"] == "table":
                add_markdown_table(document, item["headers"], item["rows"])
            else:
                add_line(document, item["text"])
    props = document.core_properties
    props.title = title
    props.subject = "Local Markdown document"
    props.author = props.last_modified_by = "OUTLANDS"
    props.comments = props.keywords = props.category = props.identifier = ""
    props.language = "en"
    props.revision = 1
    props.created = props.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def write_new_document(path: Path, data: bytes) -> None:
    if path.suffix.lower() != ".docx":
        raise BuildError("OUTPUT_EXTENSION")
    if not path.parent.is_dir():
        raise BuildError("OUTPUT_PARENT_MISSING")
    flags = os.O_WRONLY | os.O_CREAT | os.O_EXCL | getattr(os, "O_NOFOLLOW", 0)
    fd = os.open(path, flags, 0o600)
    identity = os.fstat(fd)
    try:
        with os.fdopen(fd, "wb") as stream:
            stream.write(data)
            stream.flush()
            os.fsync(stream.fileno())
    except BaseException:
        try:
            current = path.lstat()
            if (current.st_dev, current.st_ino) == (identity.st_dev, identity.st_ino):
                path.unlink()
        except OSError:
            pass
        raise


def main(argv: list[str] | None = None) -> int:
    try:
        parser = _Parser(prog="markdown_docx.py", description="Build a new DOCX from an explicit local Markdown file.")
        parser.add_argument("--input", required=True, help="UTF-8 Markdown file; no default source")
        parser.add_argument("--output", required=True, help="New .docx file; existing paths are never replaced")
        parser.add_argument("--title", default="OUTLANDS Document", help="Title, at most 160 characters")
        args = parser.parse_args(argv)
        if not _DOCX_AVAILABLE:
            raise BuildError("DEPENDENCY_UNAVAILABLE")
        if not args.title.strip() or len(args.title) > 160 or any(ord(c) < 32 for c in args.title):
            raise BuildError("TITLE_INVALID")
        source = checked_path(args.input)
        output = checked_path(args.output)
        if os.path.lexists(output):
            raise BuildError("OUTPUT_EXISTS")
        markdown = read_markdown(source)
        data = make_document(markdown, args.title.strip())
        write_new_document(output, data)
    except BuildError as error:
        print("DOCUMENT_BUILD_FAILED:" + str(error), file=sys.stderr)
        return 2
    except (OSError, UnicodeError, ValueError):
        print("DOCUMENT_BUILD_FAILED:INPUT_OR_IO", file=sys.stderr)
        return 2
    print("Document created; review its content before sharing.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
